import os
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def save_resumenes(events: list[dict], pdf_path: str, open_file: bool = False) -> str:
    events_by_date = defaultdict(list)
    for event in events:
        events_by_date[event["date"]].append(event)

    def extract_start_time(time_str: str) -> datetime:
        value = (time_str or "").strip().upper().replace(".", "")
        if "-" in value:
            value = value.split("-")[0].strip()
        if len(value) >= 3 and value[-2:] in ("AM", "PM") and not value[-3].isspace():
            value = value[:-2] + " " + value[-2:]
        try:
            return datetime.strptime(value, "%I:%M %p")
        except Exception:
            return datetime.min

    for events_day in events_by_date.values():
        events_day.sort(key=lambda e: extract_start_time(e.get("time", "")))

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"

    thin = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
    bold = Font(bold=True)
    centered = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    gray = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")

    headers = ["Time", "Organization", "Function", "Menu Style", "Room", "GTD", "Event #"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = bold
        cell.alignment = centered
        cell.border = thin

    current_row = 2
    for date, events_day in events_by_date.items():
        day_name = events_day[0].get("day", "")
        day_label = f"{day_name}, {date}"
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers))
        day_cell = ws.cell(row=current_row, column=1, value=day_label)
        day_cell.font = bold
        day_cell.alignment = centered
        day_cell.fill = gray
        day_cell.border = thin
        current_row += 1

        for event in events_day:
            ws.cell(row=current_row, column=1, value=event.get("time", ""))
            ws.cell(row=current_row, column=2, value=event.get("organization", ""))
            ws.cell(row=current_row, column=3, value=event.get("function", ""))
            ws.cell(row=current_row, column=4, value=event.get("menu_style", ""))
            ws.cell(row=current_row, column=5, value=event.get("room", ""))
            ws.cell(row=current_row, column=6, value=event.get("gtd", ""))
            ws.cell(row=current_row, column=7, value=event.get("event_order", ""))

            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=current_row, column=col)
                cell.border = thin
                cell.font = bold
                cell.alignment = centered

            current_row += 1

            raw = event.get("details", {}).get("raw", "")
            if raw:
                ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=8)
                raw_cell = ws.cell(row=current_row, column=1, value=raw)
                raw_cell.alignment = left_wrap
                raw_cell.border = thin

                lines = str(raw).count("\n") + 1
                ws.row_dimensions[current_row].height = max(20, lines * 15)
                current_row += 1

        current_row += 1

    for i in range(1, len(headers) + 1):
        col_letter = get_column_letter(i)
        max_len = 0
        for cell in ws[col_letter]:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(60, max(10, max_len + 2))

    file_name_no_ext = os.path.splitext(os.path.basename(pdf_path))[0]
    out_path = f"ready_to_print/resumenes/{file_name_no_ext}_all.xlsx"
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)

    absolute_path = os.path.abspath(out_path)
    if open_file:
        os.startfile(absolute_path)
    print(f"\nFile '{absolute_path}' generated.")
    return absolute_path
