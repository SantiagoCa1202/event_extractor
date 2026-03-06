from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment

def generate_waste_food_sheet(events, filename="events.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Waste Food"

    # Definir borde fino
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    headers = ["Item", "Pans Sent (Kitchen)", "Pans Used (Service)", "Waste"]
    row = 1

    for event in events:
        # ====== SECCIÓN ENCABEZADOS ======
        for label, key in [
            ("Event Order:", "event_order"),
            ("Time:", "time"),
            ("Room:", "room"),
            ("Function:", "function"),
            ("Date:", "date"),
            ("Day:", "day"),
            ("GTD:", "gtd"),
        ]:
            ws[f"A{row}"] = label
            value = event.get(key, "")
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)

            merged_cell = ws.cell(row=row, column=2, value=value)
            merged_cell.font = Font(size=20)
            merged_cell.alignment = Alignment(vertical="center", wrap_text=True)
            merged_cell.border = thin_border

            for col in range(3, 5):
                ws.cell(row=row, column=col).border = thin_border

            ws[f"A{row}"].font = Font(bold=True, size=20)
            ws[f"A{row}"].border = thin_border
            ws[f"A{row}"].alignment = Alignment(vertical="center")
            row += 1

        row += 1  # Espacio antes de la tabla

        # ====== TABLA DE WASTE ======
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = Font(bold=True, size=20)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        row += 1

        waste_items = event.get("details", {}).get("food", {}).get("waste_items", [])
        for item in waste_items:
            for col in range(1, 5):
                cell = ws.cell(row=row, column=col)
                cell.font = Font(bold=True, size=20)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.cell(row=row, column=1, value=item.title())
            row += 1

        row += 2  # espacio antes de siguiente evento

    # Ajustar ancho columnas A–D
    for col in range(1, 5):
        ws.column_dimensions[chr(64 + col)].width = 25

    # ====== GUARDAR ARCHIVO ======
    wb.save(filename)
    print(f"✅ Archivo Excel creado: {filename}")
