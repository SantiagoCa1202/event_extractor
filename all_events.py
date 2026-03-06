import os
import re
from collections import defaultdict
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def save_events_to_excel(events: list[dict], pdf_path: str, open_file: bool = False) -> str:
    """Process events and save them into a formatted Excel file."""

    events_by_date = defaultdict(list)
    for event in events:
        events_by_date[event["date"]].append(event)

    def extract_start_time(time_str: str) -> datetime:
        normalized = (time_str or "").strip().upper().replace(".", "")
        start_time = normalized.split("-")[0].strip() if "-" in normalized else normalized
        if len(start_time) >= 3 and start_time[-2:] in ("AM", "PM") and not start_time[-3].isspace():
            start_time = start_time[:-2] + " " + start_time[-2:]
        try:
            return datetime.strptime(start_time, "%I:%M %p")
        except Exception:
            return datetime.min

    for events_day in events_by_date.values():
        events_day.sort(key=lambda e: extract_start_time(e.get("time", "")))

    excel_rows = []
    for date, events_day in events_by_date.items():
        day_name = events_day[0]["day"]
        excel_rows.append({"": f"{day_name}, {date}"})
        for event in events_day:
            excel_rows.append(
                {
                    "Time": event["time"],
                    "Organization": event["organization"],
                    "Function": f"{event['function']}",
                    "Menu Style": event["menu_style"],
                    "Room": event["room"],
                    "GTD": event["gtd"],
                    "Event #": event["event_order"],
                }
            )
        excel_rows.append({})

    df_final = pd.DataFrame(excel_rows)

    file_name_no_ext = os.path.splitext(os.path.basename(pdf_path))[0]
    excel_path_all = f"ready_to_print/all/{file_name_no_ext}_all.xlsx"
    os.makedirs(os.path.dirname(excel_path_all), exist_ok=True)

    df_final.to_excel(excel_path_all, index=False)

    wb = load_workbook(excel_path_all)
    ws = wb.active

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    bold_font = Font(bold=True)
    centered = Alignment(horizontal="center", vertical="center")
    gray_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")

    max_col = ws.max_column

    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_length + 2

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=max_col):
        row_number = row[0].row
        for cell in row:
            cell.border = thin_border

        first_value = row[0].value
        if isinstance(first_value, str) and re.match(
            r"^[A-Za-z]+,\s+[A-Za-z]+\s+\d{1,2},\s+\d{4}$", first_value.strip()
        ):
            ws.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=max_col)
            cell = ws.cell(row=row_number, column=1)
            cell.font = bold_font
            cell.alignment = centered
            cell.fill = gray_fill

    wb.save(excel_path_all)

    absolute_path = os.path.abspath(excel_path_all)
    if open_file:
        os.startfile(absolute_path)
    print(f"\nFile '{absolute_path}' has been updated.")
    return absolute_path
